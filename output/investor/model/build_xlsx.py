"""Writes output/investor/zanjir-forecast-monthly.xlsx: the 60-month monthly forecast behind the final PDF."""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
import zanjir_model_final as M

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "zanjir-forecast-monthly.xlsx")
MONTHS = "فروردین اردیبهشت خرداد تیر مرداد شهریور مهر آبان آذر دی بهمن اسفند".split()


def jal(m):
    j = 8 + (m - 1)
    return f"{MONTHS[(j - 1) % 12]} {1405 + (j - 1) // 12}"


COLS = [
    ("ماه", None), ("ماه شمسی", None),
    ("نصب — ایران", lambda r, t: r["seg"]["ir"]["inst"][t]),
    ("نصب — گوگل‌پلی", lambda r, t: r["seg"]["gp"]["inst"][t]),
    ("نصب — اپ‌استور انگلیسی", lambda r, t: r["seg"]["ase"]["inst"][t]),
    ("نصب — اپ‌استور فارسی", lambda r, t: r["seg"]["asf"]["inst"][t]),
    ("DAU — ایران", lambda r, t: r["seg"]["ir"]["dau"][t]),
    ("DAU — گوگل‌پلی", lambda r, t: r["seg"]["gp"]["dau"][t]),
    ("DAU — اپ‌استور انگلیسی", lambda r, t: r["seg"]["ase"]["dau"][t]),
    ("DAU — اپ‌استور فارسی", lambda r, t: r["seg"]["asf"]["dau"][t]),
    ("درآمد — ایران", lambda r, t: r["seg"]["ir"]["rev"][t]),
    ("درآمد — گوگل‌پلی", lambda r, t: r["seg"]["gp"]["rev"][t]),
    ("درآمد — اپ‌استور انگلیسی", lambda r, t: r["seg"]["ase"]["rev"][t]),
    ("درآمد — اپ‌استور فارسی", lambda r, t: r["seg"]["asf"]["rev"][t]),
    ("درآمد خالص کل", lambda r, t: r["rev"][t]),
    ("هزینه — مؤسس", lambda r, t: r["cost"]["founder"][t]),
    ("هزینه — خدمات بیرونی ایران", lambda r, t: r["cost"]["outsource_ir"][t]),
    ("هزینه — سرور و پیامک ایران", lambda r, t: r["cost"]["infra_ir"][t]),
    ("هزینه — تبلیغ جذب", lambda r, t: r["cost"]["ua_ir"][t] + r["cost"]["intl_ua"][t]),
    ("هزینه — خارجی (سرور، حساب‌ها، خدمات)", lambda r, t: r["cost"]["intl_host"][t] + r["cost"]["intl_outsource"][t] + r["cost"]["intl_accounts"][t]),
    ("هزینه — یک‌باره", lambda r, t: r["cost"]["one_off"][t]),
    ("هزینهٔ کل", lambda r, t: r["cost"]["total"][t]),
    ("سود (زیان) عملیاتی", lambda r, t: r["ebit"][t]),
    ("مالیات بر درآمد (تعهدی)", lambda r, t: r["tax"][t]),
    ("سود (زیان) خالص", lambda r, t: r["net"][t]),
    ("جریان نقد (با تأخیر دریافت)", lambda r, t: r["cash"][t]),
    ("مانده نقدی تجمعی", lambda r, t: r["cum"][t]),
]

wb = Workbook()
info = wb.active
info.title = "راهنما"
info.sheet_view.rightToLeft = True
for row in [
    ["زنجیر — پیش‌بینی ماهانهٔ ۶۰ ماهه (پشتوانهٔ سند «طرح درآمدی و پیش‌بینی مالی»، مهر ۱۴۰۵)"],
    ["مبالغ: میلیون تومان، قیمت ثابت مهر ۱۴۰۵؛ هر دلار ۲۶۶٬۰۰۰ تومان. نصب و DAU: نفر."],
    ["ماه ۱ = آبان ۱۴۰۵. ماه‌های ۳۷ تا ۶۰ برون‌یابی است."],
    ["درآمدها خالص و پس از همهٔ کسورات بخش ۲ سند (کارمزد فروشگاه، مالیات بر ارزش افزوده، تورم، قطعی، سهم و مالیات همکار، انتقال) است."],
    ["جریان نقد: درآمد ایران ۱ ماه و درآمد خارجی ۲ ماه دیرتر؛ مالیات سال هر سال در ماه چهارم سال بعد پرداخت می‌شود."],
    ["اقساط سرمایه (۱٬۷۰۰ در ماه ۱، ۱٬۱۰۰ در ماه ۷، ۱٬۰۰۰ در ماه ۱۹) در «مانده نقدی تجمعی» نیامده‌اند."],
]:
    info.append(row)
info.column_dimensions["A"].width = 120

head_fill = PatternFill("solid", fgColor="1E1B2E")
for s in M.SCEN:
    r = M.scenario(s)
    ws = wb.create_sheet(M.SCEN_FA[s])
    ws.sheet_view.rightToLeft = True
    ws.append([c for c, _ in COLS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 42
    for t in range(M.T):
        ws.append([t + 1, jal(t + 1)] + [round(float(f(r, t)), 1) for _, f in COLS[2:]])
    for i in range(1, len(COLS) + 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 14
    for row in ws.iter_rows(min_row=2, min_col=3):
        for cell in row:
            cell.number_format = '#,##0;(#,##0)'
    ws.freeze_panes = "C2"

wb.save(OUT)
print("wrote", OUT)
